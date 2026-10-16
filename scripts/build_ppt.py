"""기능·구현 기술 설명 PPT 생성 (python-pptx) → docs/발표자료.pptx

사용: python scripts/build_ppt.py [--skip-tests]
- 실모델 캡처(파일명에 _mock 없음)가 있으면 우선 사용, 없으면 _mock 캡처를 쓰고 슬라이드에 'mock 캡처'로 표기
- 기본으로 pytest를 실제 실행해 그 결과 요약 줄을 '검증 결과' 슬라이드에 그대로 넣음
- 실제 구현된 내용만 기술함. 샘플 데이터 수치는 인용하지 않음
"""
from __future__ import annotations

import argparse
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
SHOTS = ROOT / "docs" / "screenshots"
OUT = ROOT / "docs" / "발표자료.pptx"

FONT = "Malgun Gothic"
FOREST = RGBColor(0x1F, 0x3A, 0x2B)
MOSS = RGBColor(0x5E, 0x8C, 0x4A)
TINT = RGBColor(0xEE, 0xF3, 0xEA)
ACCENT = RGBColor(0xD9, 0x6C, 0x2C)
INK = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x5F, 0x63, 0x5C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
WARN_BG = RGBColor(0xFD, 0xF1, 0xE6)

W, H = Inches(13.333), Inches(7.5)
TMP = Path(tempfile.mkdtemp(prefix="ppt_"))


# ---------------------------------------------------------------- 공통 도우미
def _font(run, size, bold=False, color=INK, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", FONT)


def text(slide, x, y, w, h, lines, size=16, color=INK, bold=False, bullet=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, space_after=6):
    """lines: str | list[str | (str, dict)]"""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    items = [lines] if isinstance(lines, str) else lines
    for i, item in enumerate(items):
        content, opt = (item, {}) if isinstance(item, str) else item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        r = p.add_run()
        r.text = (("• " if bullet and not opt.get("nobullet") else "") + content)
        _font(r, opt.get("size", size), opt.get("bold", bold), opt.get("color", color), opt.get("italic", False))
    return box


def rect(slide, x, y, w, h, fill=TINT, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.shadow.inherit = False
    return s


def box_label(slide, x, y, w, h, title, body=None, fill=TINT, title_color=FOREST, body_color=INK, size=14):
    rect(slide, x, y, w, h, fill=fill)
    lines = [(title, {"bold": True, "color": title_color, "size": size + 1})]
    if body:
        lines += [(b, {"size": size - 1, "color": body_color}) for b in body]
    text(slide, x + 0.12, y + 0.08, w - 0.24, h - 0.16, lines, space_after=3)


def arrow(slide, x1, y1, x2, y2, color=MUTED):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(1.75)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return c


def new_slide(prs, title=None, subtitle=None, dark=False):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = FOREST if dark else WHITE
    if title:
        text(s, 0.6, 0.4, 12.1, 0.8, title, size=32, bold=True, color=WHITE if dark else FOREST)
    if subtitle:
        text(s, 0.6, 1.12, 12.1, 0.5, subtitle, size=15, color=TINT if dark else MUTED)
    return s


def pick_shot(stem: str) -> tuple[Path | None, bool]:
    real = SHOTS / f"{stem}.png"
    mock = SHOTS / f"{stem}_mock.png"
    if real.exists():
        return real, False
    if mock.exists():
        return mock, True
    return None, False


def crop_to(path: Path, aspect: float, top: int = 0, left: int = 0) -> Path:
    img = Image.open(path)
    w = img.width - left
    h = min(int(w / aspect), img.height - top)
    out = TMP / f"{path.stem}_{top}_{left}.png"
    img.crop((left, top, left + w, top + h)).save(out)
    return out


def screenshot(slide, stem, x, y, w, h, top=0, left=0):
    path, is_mock = pick_shot(stem)
    if path is None:
        rect(slide, x, y, w, h, fill=TINT)
        text(slide, x, y + h / 2 - 0.3, w, 0.6, f"캡처 없음: {stem} (scripts/capture_screens.py 실행 필요)",
             size=13, color=MUTED, align=PP_ALIGN.CENTER)
        return
    img = crop_to(path, w / h, top, left)
    pic = slide.shapes.add_picture(str(img), Inches(x), Inches(y), Inches(w), Inches(h))
    pic.line.color.rgb = RGBColor(0xC9, 0xD3, 0xC4)
    pic.line.width = Pt(1)
    label = "mock 모드 캡처 · 실모델 결과 아님" if is_mock else "실모델 캡처"
    text(slide, x, y + h + 0.05, w, 0.35, f"{path.name} · {label}", size=10, color=ACCENT if is_mock else MUTED)


def table(slide, x, y, w, rows, widths, size=12, row_h=0.42, accent_col=None, header_fill=FOREST):
    """rows[0]은 헤더. 셀 값의 '\n'은 줄바꿈. widths 합은 w와 같게 지정."""
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
    tbl = shape.table
    for j, wv in enumerate(widths):
        tbl.columns[j].width = Inches(wv)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if i == 0 else (TINT if i % 2 else WHITE)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.text_frame.word_wrap = True
            for k, line in enumerate(str(val).split("\n")):
                para = cell.text_frame.paragraphs[0] if k == 0 else cell.text_frame.add_paragraph()
                r = para.add_run()
                r.text = line
                color = WHITE if i == 0 else (ACCENT if accent_col is not None and j == accent_col else INK)
                _font(r, size, bold=(i == 0 or j == 0), color=color)
    return tbl


def fit_shot(slide, stem, x, y, w, h, caption=True):
    """비율을 유지해 박스 안에 배치. 박스보다 세로로 긴 이미지는 위쪽부터 박스 비율로 잘라 씀."""
    path, is_mock = pick_shot(stem)
    if path is None:
        rect(slide, x, y, w, h, fill=TINT)
        text(slide, x, y + h / 2 - 0.3, w, 0.6, f"캡처 없음: {stem}", size=12, color=MUTED, align=PP_ALIGN.CENTER)
        return
    img = Image.open(path)
    box_ratio = h / w
    if img.height / img.width > box_ratio:  # 세로로 길면 위쪽을 잘라 사용
        src = crop_to(path, w / h)
        pw, ph, px, py = w, h, x, y
    else:  # 가로로 길면 폭에 맞춤
        src = path
        pw, ph = w, w * img.height / img.width
        px, py = x, y  # 위쪽 정렬
    pic = slide.shapes.add_picture(str(src), Inches(px), Inches(py), Inches(pw), Inches(ph))
    pic.line.color.rgb = RGBColor(0xC9, 0xD3, 0xC4)
    pic.line.width = Pt(1)
    if caption:
        label = "mock 캡처" if is_mock else "실모델 캡처"
        text(slide, px, py + ph + 0.03, pw, 0.3, f"{path.stem} · {label}", size=10, color=ACCENT if is_mock else MUTED)


TEST_DESC = {
    "test_llm": "factory·mock tool_calls 형식, 429 백오프",
    "test_data": "샘플 대체, KAMIS 컬럼 매핑, KOSIS wide→long",
    "test_rag": "출처 반환, 문서 근거 없음, e5 접두어, 실 Embedding 분리도",
    "test_vision": "JSON 스키마, 캐시, 1회 재요청, analysis_logs",
    "test_agent": "질의별 Tool 선택 4건, 6회 제한, 오류 처리",
    "test_zip": "제출 zip 포함·제외 규칙",
}


def run_tests(skip: bool) -> tuple[str, list[tuple]]:
    """(요약 줄, 모듈별 [모듈, 설명, 통과, 실패, skip]) — pytest를 실제 실행한 결과."""
    if skip:
        return "테스트 미실행 (--skip-tests)", []
    import os
    import xml.etree.ElementTree as ET

    junit = TMP / "junit.xml"
    proc = subprocess.run([sys.executable, "-m", "pytest", "-q", "-p", "no:cacheprovider", f"--junitxml={junit}"],
                          cwd=ROOT, capture_output=True, text=True, env={**os.environ, "LLM_PROVIDER": "mock"})
    lines = [l for l in proc.stdout.strip().splitlines() if re.search(r"\d+ (passed|failed)", l)]
    summary = lines[-1].strip("= ") if lines else f"pytest 실행 결과 해석 실패 (exit {proc.returncode})"
    counts: dict[str, list[int]] = {}
    if junit.exists():
        for case in ET.parse(junit).getroot().iter("testcase"):
            mod = case.get("classname", "").split(".")[-1]
            c = counts.setdefault(mod, [0, 0, 0])
            if case.find("failure") is not None or case.find("error") is not None:
                c[1] += 1
            elif case.find("skipped") is not None:
                c[2] += 1
            else:
                c[0] += 1
    rows = [(m, TEST_DESC.get(m, ""), *c) for m, c in sorted(counts.items())]
    return summary, rows


RAG_QUERIES = [("적정 생육 온도", True), ("푸른곰팡이병 방제", True), ("환기 관리", True),
               ("버섯파리 방제 방법", True), ("자동차 엔진오일 교체 주기", False), ("비트코인 시세 전망", False),
               ("오늘 날씨 어때", False)]


def measure_rag() -> tuple[str, float, list[tuple]]:
    """(Embedding 백엔드, 임계값, [질의, 구분, top-1 유사도, 출처, 판정]) — 빌드 시 실제 검색 결과."""
    from src.rag import retriever
    from src.rag.embeddings import default_threshold, get_embeddings

    _, backend, _ = get_embeddings()
    thr = default_threshold(backend)
    rows = []
    for q, relevant in RAG_QUERIES:
        hit = retriever.search(q, threshold=0.0)["hits"][0]
        passed = hit["score"] >= thr
        verdict = ("근거 제시" if passed else "문서 근거 없음")
        ok = passed == relevant
        rows.append((q, "관련" if relevant else "무관", f"{hit['score']:.3f}", hit["citation"],
                     verdict + (" (기대와 일치)" if ok else " (기대와 불일치)")))
    return backend, thr, rows


# ---------------------------------------------------------------- 슬라이드
def s_title(prs):
    s = new_slide(prs, dark=True)
    text(s, 0.8, 2.0, 11.5, 1.2, "멀티모달 AI Agent 기반 스마트 농업 서비스", size=40, bold=True, color=WHITE)
    text(s, 0.8, 3.15, 11.5, 0.7, "대상 작물: 느타리버섯 · Streamlit + LangGraph + VLM + RAG(Chroma)", size=20, color=TINT)
    rect(s, 0.8, 4.3, 5.6, 1.5, fill=RGBColor(0x2C, 0x4F, 0x3B))
    text(s, 1.0, 4.42, 5.3, 1.3, [("구현 범위", {"bold": True, "color": WHITE, "size": 15}),
                                  ("대시보드 · 이미지 분석 · AI 상담 · 영농일지 (4개 탭)", {"color": TINT, "size": 14}),
                                  ("Agent Tool 4종: rag_search / analyze / diary / price", {"color": TINT, "size": 14})])


def s_overview(prs):
    s = new_slide(prs, "과제 개요", "요구사항 6개 항목과 구현 위치")
    items = [
        ("1  데이터 선정", "가격(KAMIS)·지역 생산(KOSIS)·재배 매뉴얼·이미지\ndata/raw 우선, 없으면 SAMPLE"),
        ("2  대시보드", "Streamlit + Plotly 차트 4종\nKPI 3종 · 기간/지역 필터 · 출처 표기"),
        ("3  멀티모달·RAG", "VLM Zero-shot 이미지 분석(JSON)\n매뉴얼·분석결과 → Embedding → Chroma"),
        ("4  AI Agent", "LangGraph ToolNode 그래프\nTool 4종 자동 선택 · 최대 6회 반복"),
        ("5  UI", "이미지 업로드·질문 입력·대시보드\n검색 근거·Tool 이력·영농일지 출력"),
        ("6  제출물", "소스·캡처(Playwright)·PPT(python-pptx)\n제출 zip(scripts/make_zip.py)"),
    ]
    for i, (t, b) in enumerate(items):
        col, row = i % 3, i // 3
        x, y = 0.6 + col * 4.1, 1.8 + row * 2.6
        rect(s, x, y, 3.85, 2.3, fill=TINT)
        text(s, x + 0.2, y + 0.18, 3.5, 0.5, t, size=19, bold=True, color=FOREST)
        text(s, x + 0.2, y + 0.85, 3.5, 1.4, b.split("\n"), size=14, color=INK, space_after=4)


def s_data(prs):
    s = new_slide(prs, "데이터", "실데이터 투입 경로와 대체(SAMPLE) 규칙")
    rows = [
        ("구분", "투입 경로", "처리", "현재 상태"),
        ("가격", "data/raw/price/*.xlsx|csv\n(KAMIS 웹 다운로드)", "컬럼 매핑 dict, 실패 시 필요 컬럼 UI 표시\nKAMIS API는 키+품목코드 있을 때만", "SAMPLE"),
        ("지역 생산", "data/raw/production/*.csv\n(KOSIS 특용작물생산실적)", "wide→long, cp949/utf-8 자동 판별\n전국·계 행 제외", "SAMPLE"),
        ("재배 매뉴얼", "data/raw/docs/*.pdf|md", "PyPDFLoader, 파일명·페이지 메타데이터", "샘플 매뉴얼(일반 상식)"),
        ("이미지", "data/raw/images/*.jpg|png", "시연·테스트용(학습 없음, Zero-shot)", "mock 더미 이미지"),
    ]
    tbl = s.shapes.add_table(len(rows), 4, Inches(0.6), Inches(1.8), Inches(12.1), Inches(4.2)).table
    widths = [1.6, 3.6, 4.9, 2.0]
    for j, wv in enumerate(widths):
        tbl.columns[j].width = Inches(wv)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = FOREST if i == 0 else (TINT if i % 2 else WHITE)
            cell.text_frame.word_wrap = True
            for k, line in enumerate(val.split("\n")):
                p = cell.text_frame.paragraphs[0] if k == 0 else cell.text_frame.add_paragraph()
                r = p.add_run()
                r.text = line
                _font(r, 13 if i else 14, bold=(i == 0 or j == 0), color=WHITE if i == 0 else
                      (ACCENT if j == 3 else INK))
    rect(s, 0.6, 6.25, 12.1, 0.7, fill=WARN_BG)
    text(s, 0.8, 6.32, 11.8, 0.6, "샘플 가격·생산량은 합성값(is_sample=True)이며 실측 통계가 아님. 대시보드 상단과 각 차트 하단에 SAMPLE로 표기함",
         size=14, color=ACCENT, bold=True, anchor=MSO_ANCHOR.MIDDLE)


def s_arch(prs):
    s = new_slide(prs, "시스템 아키텍처", "공급자 생성은 src/llm/factory.py 한 곳에서만 수행")
    # UI
    box_label(s, 0.6, 1.75, 12.1, 0.95, "Streamlit UI (app.py)",
              ["📊 대시보드   🔬 이미지 분석   💬 AI 상담   📝 영농일지   · 사이드바: LLM_PROVIDER·모델·Embedding 표시"],
              fill=FOREST, title_color=WHITE, body_color=TINT)
    # 중간: Agent / Tools / factory
    box_label(s, 0.6, 3.15, 3.6, 1.6, "LangGraph Agent", ["agent 노드 ↔ ToolNode", "최대 6회 · 429/연결 오류 백오프 2회"])
    box_label(s, 4.6, 3.15, 4.3, 1.6, "Tools (src/agent/tools.py)",
              ["rag_search · analyze_mushroom_image", "write_farm_diary · get_mushroom_price"])
    box_label(s, 9.3, 3.15, 3.4, 1.6, "LLM factory", ["mock (기본, 키 불필요)", "gemini · ollama", "LLM_MODEL / VLM_MODEL 환경변수"])
    # 하단 계층
    bottom = [
        ("Chroma", ["manuals", "analysis_logs"]),
        ("Embedding", ["sentence-transformers", "e5 접두어 · 해싱 폴백"]),
        ("Vision", ["Pydantic JSON 스키마", "SHA-256 캐시"]),
        ("Data loaders", ["data/raw → SAMPLE", "KAMIS API(선택)"]),
        ("outputs/", ["diary/*.md", "cache · uploads"]),
    ]
    for i, (t, b) in enumerate(bottom):
        box_label(s, 0.6 + i * 2.45, 5.35, 2.25, 1.35, t, b, size=13)
    arrow(s, 2.4, 2.7, 2.4, 3.15)
    arrow(s, 4.2, 3.95, 4.6, 3.95)
    arrow(s, 9.3, 3.95, 8.9, 3.95)
    for x in (1.7, 4.15, 6.6, 9.05, 11.5):  # Tools → 하단 계층
        arrow(s, min(max(x, 4.9), 8.6), 4.75, x, 5.35)


def s_dashboard(prs):
    s = new_slide(prs, "대시보드 (탭 1)", "Streamlit + Plotly · 데이터 출처를 각 차트 하단에 표기")
    screenshot(s, "01_dashboard", 0.6, 1.7, 7.4, 5.2)
    text(s, 8.3, 1.8, 4.5, 5.0, [
        ("차트 4종", {"bold": True, "color": FOREST, "size": 17}),
        "① 가격 시계열 라인 — 일/월 전환, 도매·소매 구분",
        "② 시도별 생산량 가로 막대 — 연도 선택",
        "③ 시도별 생산량 전년 대비 증감률 히트맵",
        "④ 도매 가격 월별 계절성 히트맵 (연 × 월)",
        ("KPI · 필터", {"bold": True, "color": FOREST, "size": 17}),
        "최근 도매가 · 전월 대비 %(월평균) · 최대 생산 지역",
        "기간(date range) · 지역(multiselect)",
        ("이중 축 미사용, 범주색 고정 순서", {"color": MUTED, "size": 13}),
    ], size=14, space_after=7)


def s_vision(prs):
    s = new_slide(prs, "이미지 분석 (VLM Zero-shot)", "src/vision/analyzer.py · 학습 코드 없음(학습 데이터 없음)")
    screenshot(s, "02_image_analysis", 0.6, 1.7, 6.6, 4.9, left=300)
    box_label(s, 7.5, 1.75, 5.2, 2.35, "구조화 출력 스키마 (Pydantic)", [
        "growth_stage: 균사생장/원기형성/자실체생장/수확적기/과숙/판단불가",
        "disease_suspected: 정상/푸른곰팡이병/세균성갈색무늬병/버섯파리 피해/기타/판단불가",
        "confidence 0~1 · visual_evidence[] · recommended_actions[] · summary_ko",
    ], size=13)
    box_label(s, 7.5, 4.3, 5.2, 2.3, "처리 규칙", [
        "프롬프트: 관찰 근거만 기술, 불확실하면 판단불가·낮은 신뢰도",
        "파싱 실패 → 1회 재요청 → 실패 시 판단불가 반환(캐시 안 함)",
        "SHA-256(+공급자·모델) 키 캐시: 같은 이미지는 재호출 없음",
        "결과를 Chroma analysis_logs에 메타데이터와 함께 저장",
    ], size=13)


def s_rag(prs):
    s = new_slide(prs, "RAG: 재배 매뉴얼 검색", "src/rag/ingest.py · retriever.py · python -m src.rag.ingest 로 재색인")
    steps = [("문서 로딩", "pdf(PyPDFLoader)·md\n파일명·페이지 유지"), ("청킹", "Recursive splitter\n600자 / overlap 100"),
             ("Embedding", "multilingual-e5-small\nquery:/passage: 접두어"), ("Chroma", "컬렉션 manuals\ncosine 거리"),
             ("검색·판정", "top-k 4, 임계값 미만\n→ '문서 근거 없음'")]
    for i, (t, b) in enumerate(steps):
        x = 0.6 + i * 2.5
        rect(s, x, 1.9, 2.15, 1.9, fill=FOREST if i == 3 else TINT)
        text(s, x + 0.12, 2.0, 1.95, 1.75, [(t, {"bold": True, "size": 16, "color": WHITE if i == 3 else FOREST})] +
             [(l, {"size": 13, "color": TINT if i == 3 else INK}) for l in b.split("\n")], space_after=4)
        if i < 4:
            arrow(s, x + 2.15, 2.85, x + 2.5, 2.85)
    box_label(s, 0.6, 4.2, 6.0, 2.6, "출처 표기", [
        "모든 검색 결과에 파일명 · 페이지(PDF) · 청크 id · 유사도 포함",
        "Agent 답변 끝에 [번호]와 출처 인용 (시스템 프롬프트 규칙)",
        "UI: Tool 호출 이력 expander에 검색 근거 표 출력",
    ], size=14)
    box_label(s, 6.9, 4.2, 5.8, 2.6, "운영 처리", [
        "앱 최초 실행 시 vectorstore가 비면 자동 색인",
        "EMBEDDING_MODEL=BAAI/bge-m3 전환 가능 · 기본 임계값 0.855(e5-small 보정값)",
        "모델 로드 실패 시 해싱 폴백 + 사이드바 경고, 백엔드 변경 시 자동 재색인",
    ], fill=WARN_BG, title_color=ACCENT, size=14)


def s_agent(prs):
    s = new_slide(prs, "Agent · Tool 흐름", "src/agent/graph.py · StateGraph(agent ↔ ToolNode)")
    y0 = 2.15
    box_label(s, 0.6, y0, 2.6, 1.2, "사용자 입력", ["질문 + [첨부 이미지 경로]"])
    box_label(s, 3.8, y0, 2.8, 1.2, "agent 노드", ["LLM.bind_tools → tool_calls"], fill=FOREST, title_color=WHITE, body_color=TINT)
    box_label(s, 7.2, y0, 2.6, 1.2, "ToolNode", ["Tool 실행 → ToolMessage"])
    box_label(s, 10.4, y0, 2.3, 1.2, "최종 답변", ["출처·주의사항 포함"])
    arrow(s, 3.2, y0 + 0.6, 3.8, y0 + 0.6)
    arrow(s, 6.6, y0 + 0.4, 7.2, y0 + 0.4)
    arrow(s, 7.2, y0 + 0.85, 6.6, y0 + 0.85)
    # 종료 경로: agent 위로 우회해 최종 답변으로 (박스를 가로지르지 않게)
    for x1, y1, x2, y2 in ((5.2, y0, 5.2, 1.8), (5.2, 1.8, 11.55, 1.8)):
        ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        ln.line.color.rgb = MOSS
        ln.line.width = Pt(1.75)
    arrow(s, 11.55, 1.8, 11.55, y0, color=MOSS)
    text(s, 7.6, 1.42, 4.0, 0.35, "tool_calls 없음 또는 6회 도달 → 종료", size=12, color=MOSS)
    tools = [
        ("rag_search(query)", "재배 방법·환경·방제 질문, 문서 요약"),
        ("analyze_mushroom_image(image_path, question)", "첨부 이미지 상태 질문"),
        ("write_farm_diary(date, notes)", "\"오늘 일지 써줘\" — 메모의 온도·습도·CO2 반영"),
        ("get_mushroom_price(period)", "가격·시세 질문 — 1w/1m/3m/6m/1y/all"),
    ]
    for i, (t, b) in enumerate(tools):
        col, row = i % 2, i // 2
        box_label(s, 0.6 + col * 6.15, 3.75 + row * 1.3, 5.95, 1.1, t, [b], size=13)
    text(s, 0.6, 6.45, 12.1, 0.6, "반복 6회 제한: 마지막 회차는 Tool 없이 답변 생성 · 429/연결 오류는 지수 백오프 2회 후 'API 한도 초과 또는 모델 서버 미연결' 표시",
         size=13, color=MUTED)


def s_demo(prs):
    s = new_slide(prs, "시연 시나리오", "LLM_PROVIDER=mock으로 키 없이 전체 흐름 재현 가능")
    steps = [
        ("대시보드", "SAMPLE 배지 확인 → 기간·지역 필터, 일/월 전환"),
        ("이미지 분석", "이미지 업로드 → 분석 실행 → JSON·캐시 적중 확인(재실행)"),
        ("문서 질의", "\"자실체 생육기 적정 온도와 환기 관리는?\" → rag_search, 출처 표"),
        ("이미지 질의", "AI 상담에 이미지 첨부 + \"상태 어때?\" → analyze_mushroom_image"),
        ("가격 질의", "\"최근 가격 알려줘\" → get_mushroom_price(1m)"),
        ("영농일지", "\"오늘 일지 써줘. 온도 16도 습도 90%\" 또는 탭 4에서 작성 → md 다운로드"),
    ]
    for i, (t, b) in enumerate(steps):
        y = 1.8 + i * 0.85
        c = rect(s, 0.6, y, 0.62, 0.62, fill=FOREST, shape=MSO_SHAPE.OVAL)
        text(s, 0.6, y + 0.06, 0.62, 0.5, str(i + 1), size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        text(s, 1.45, y + 0.08, 2.4, 0.5, t, size=17, bold=True, color=FOREST)
        text(s, 3.9, y + 0.1, 8.8, 0.5, b, size=15)


def s_requirements(prs, test_rows):
    s = new_slide(prs, "요구사항 ↔ 구현 ↔ 검증", "과제 요구사항 원문 기준 대응표 · 검증은 빌드 시 pytest 실행 결과")
    by_mod = {r[0]: r for r in test_rows}

    def res(*mods):
        got = [by_mod[m] for m in mods if m in by_mod]
        if not got:
            return "미실행"
        p, f = sum(g[2] for g in got), sum(g[3] for g in got)
        return f"{p}건 통과" + (f" / {f}건 실패" if f else "")

    rows = [
        ("요구사항", "구현 내용", "구현 파일", "검증"),
        ("1. 데이터 선정", "가격·지역 생산·매뉴얼·이미지, data/raw 우선 → SAMPLE 대체", "src/data/loaders.py, kamis.py", res("test_data")),
        ("2. 대시보드", "Plotly 차트 4종, KPI 3종, 기간·지역 필터, 출처 표기", "src/dashboard/charts.py, app.py", "Playwright 캡처 · 예외 0건"),
        ("3. 이미지 분석", "VLM Zero-shot, Pydantic JSON, SHA-256 캐시", "src/vision/analyzer.py", res("test_vision")),
        ("3. RAG", "청킹 → e5 Embedding → Chroma(manuals, analysis_logs)", "src/rag/*.py", res("test_rag")),
        ("4. AI Agent", "LangGraph agent↔ToolNode, Tool 4종, 최대 6회", "src/agent/graph.py, tools.py", res("test_agent", "test_llm")),
        ("5. UI", "업로드·질문·대시보드·검색 근거·Tool 이력·영농일지", "app.py", f"Playwright 캡처 {len(list(SHOTS.glob('*.png')))}장"),
        ("6. 제출물", "캡처·PPT·zip 자동 생성", "scripts/*.py", res("test_zip")),
    ]
    table(s, 0.6, 1.75, 12.1, rows, [1.75, 5.0, 3.25, 2.1], size=12, row_h=0.6, accent_col=3)


def s_stack(prs):
    s = new_slide(prs, "기술 스택과 선정 근거", "공급자·Embedding은 환경변수로 전환 (코드 수정 없음)")
    rows = [
        ("구성요소", "적용 기술", "선정 근거", "대안 · 전환 방법"),
        ("UI · 시각화", "Streamlit, Plotly", "Python 단일 언어로 탭·업로드·차트 구현", "—"),
        ("Agent", "LangGraph StateGraph + ToolNode", "반복 횟수 제어·Tool 이력 추출이 명시적", "create_react_agent"),
        ("LLM · VLM", "mock / Gemini / Ollama", "키 없이 개발·테스트(mock), 실모델은 env로 전환", "LLM_PROVIDER, LLM_MODEL, VLM_MODEL"),
        ("Embedding", "sentence-transformers\nmultilingual-e5-small", "로컬 실행, API 키 불필요, 한국어 지원", "EMBEDDING_MODEL=BAAI/bge-m3 · hash 폴백"),
        ("Vector DB", "Chroma (persist: vectorstore/)", "파일 기반 영속화, 별도 서버 불필요", "—"),
        ("구조화 출력", "Pydantic v2 스키마 검증", "허용값 외 응답 차단, 실패 시 1회 재요청", "—"),
        ("데이터", "pandas, openpyxl, pypdf", "KAMIS xlsx·KOSIS csv(cp949)·PDF 로딩", "KAMIS Open API(선택)"),
        ("산출물", "Playwright, python-pptx, pytest", "캡처·PPT·테스트를 스크립트로 재현", "—"),
    ]
    table(s, 0.6, 1.75, 12.1, rows, [1.8, 3.2, 4.1, 3.0], size=12, row_h=0.56)


def s_dataflow(prs):
    s = new_slide(prs, "데이터 흐름", "입력 데이터별 처리 경로와 저장 위치")
    lanes = [
        ("가격·생산 통계", "data/raw/price · production\n(없으면 data/sample)", "loaders.py\n컬럼 매핑 · wide→long", "대시보드 차트·KPI\nget_mushroom_price"),
        ("재배 매뉴얼", "data/raw/docs (pdf·md)\n(없으면 샘플 매뉴얼)", "ingest.py\n청킹 600/100 → e5", "Chroma manuals\n→ rag_search"),
        ("재배 이미지", "업로드 → outputs/uploads\n(시연: data/raw/images)", "analyzer.py\nVLM → Pydantic JSON", "outputs/cache · analysis_logs\n→ 분석 탭·일지"),
        ("농가 입력", "온도·습도·CO2·작업 메모\n(탭 4 또는 상담 메시지)", "writer.py\n고정 양식 + 익일 계획", "outputs/diary/\n날짜별 Markdown → 다운로드"),
    ]
    heads = ["입력", "저장 · 대체 규칙", "처리 모듈", "결과 · 사용처"]
    xs, ws = [0.6, 3.05, 6.35, 9.65], [2.15, 2.95, 2.95, 3.05]
    for j, hd in enumerate(heads):
        text(s, xs[j], 1.65, ws[j], 0.4, hd, size=14, bold=True, color=MUTED)
    for i, lane in enumerate(lanes):
        y = 2.1 + i * 1.22
        for j, val in enumerate(lane):
            fill = FOREST if j == 0 else (TINT if j != 2 else WHITE)
            rect(s, xs[j], y, ws[j], 1.0, fill=fill, line=None if j != 2 else MOSS)
            lines = val.split("\n")
            text(s, xs[j] + 0.1, y + 0.08, ws[j] - 0.2, 0.9,
                 [(lines[0], {"bold": True, "size": 14, "color": WHITE if j == 0 else FOREST})] +
                 [(l, {"size": 12, "color": INK}) for l in lines[1:]], space_after=2, anchor=MSO_ANCHOR.MIDDLE)
            if j < 3:
                arrow(s, xs[j] + ws[j], y + 0.5, xs[j + 1], y + 0.5)


def s_dashboard_detail1(prs):
    s = new_slide(prs, "대시보드 상세 ① KPI · 가격 시계열", "도매/소매 구분, 일·월 단위 전환 (일 단위 → 월평균 리샘플링)")
    fit_shot(s, "11_dashboard_kpi", 0.6, 1.65, 6.0, 1.55)
    rows = [("KPI", "산출 방식"), ("최근 도매가", "필터 기간 내 마지막 조사일 도매가"),
            ("전월 대비", "도매 월평균(MS) 최근 2개월 증감률"), ("최대 생산 지역", "최근 연도 시도별 생산량 최댓값")]
    table(s, 6.9, 1.65, 5.8, rows, [1.7, 4.1], size=12, row_h=0.39)
    fit_shot(s, "12_chart_price_daily", 0.6, 3.65, 6.0, 2.85)
    fit_shot(s, "16_chart_price_monthly", 6.75, 3.65, 6.0, 2.85)
    text(s, 0.6, 6.95, 12.1, 0.35, "모든 수치는 SAMPLE 합성값(실측 아님) · 차트 하단과 대시보드 상단에 출처 표기", size=11, color=ACCENT)


def s_dashboard_detail2(prs):
    s = new_slide(prs, "대시보드 상세 ② 생산 통계 · 계절성", "범주색 고정 순서, 크기는 단일 hue 순차 램프, 증감률은 발산 팔레트(중립 회색)")
    fit_shot(s, "13_chart_region", 0.6, 1.65, 3.9, 4.25)
    fit_shot(s, "14_chart_yoy", 4.7, 1.65, 3.9, 4.25)
    fit_shot(s, "15_chart_seasonality", 8.8, 1.65, 3.9, 1.6)
    text(s, 8.8, 3.75, 3.9, 2.4, [
        ("차트 구성", {"bold": True, "color": FOREST, "size": 15}),
        "② 시도별 생산량: 연도 선택, 가로 막대",
        "③ 전년 대비 증감률: 시도 × 연도 히트맵",
        "④ 계절성: 연 × 월 도매 월평균 히트맵",
        ("이중 축 미사용 · 지역 필터 연동", {"color": MUTED, "size": 12}),
    ], size=13, space_after=5)
    text(s, 0.6, 6.4, 12.1, 0.35, "SAMPLE 합성값이므로 지역 순위·증감률을 실제 통계로 해석하면 안 됨", size=11, color=ACCENT)


def s_vision_flow(prs):
    s = new_slide(prs, "이미지 분석 처리 흐름과 출력 스키마", "같은 이미지·같은 모델은 캐시로 재호출하지 않음")
    steps = [("이미지", "업로드 저장\nSHA-256 계산"), ("캐시 조회", "sha + 공급자 + 모델\n적중 시 즉시 반환"),
             ("VLM 호출", "base64 JPEG ≤1024px\n관찰 근거만 기술"), ("검증", "Pydantic 파싱\n실패 → 1회 재요청"),
             ("저장", "outputs/cache\nChroma analysis_logs")]
    for i, (t, b) in enumerate(steps):
        x = 0.6 + i * 2.5
        rect(s, x, 1.7, 2.15, 1.35, fill=FOREST if i == 2 else TINT)
        text(s, x + 0.1, 1.76, 1.95, 1.25, [(t, {"bold": True, "size": 15, "color": WHITE if i == 2 else FOREST})] +
             [(l, {"size": 12, "color": TINT if i == 2 else INK}) for l in b.split("\n")], space_after=2)
        if i < 4:
            arrow(s, x + 2.15, 2.37, x + 2.5, 2.37)
    rows = [
        ("필드", "타입", "허용값 · 규칙"),
        ("growth_stage", "enum", "균사생장 / 원기형성 / 자실체생장 / 수확적기 / 과숙 / 판단불가"),
        ("disease_suspected", "enum", "정상 / 푸른곰팡이병 / 세균성갈색무늬병 / 버섯파리 피해 / 기타 / 판단불가"),
        ("confidence", "float", "0.0 ~ 1.0 (범위 밖이면 검증 실패), 불확실하면 ≤ 0.3"),
        ("visual_evidence", "list[str]", "이미지에서 관찰된 근거만"),
        ("recommended_actions", "list[str]", "권장 조치"),
        ("summary_ko", "str", "한국어 요약 · 2회 파싱 실패 시 판단불가 객체 반환(캐시 안 함)"),
    ]
    table(s, 0.6, 3.4, 7.3, rows, [2.0, 1.0, 4.3], size=11, row_h=0.47)
    fit_shot(s, "21_vision_result", 8.15, 3.4, 4.55, 3.3)


def s_rag_eval(prs, rag):
    backend, thr, rows = rag
    s = new_slide(prs, "RAG 검색 실측", f"빌드 시 실제 검색 결과 · Embedding: {backend} · 임계값 {thr}")
    table(s, 0.6, 1.75, 12.1, [("질의", "구분", "top-1 유사도", "top-1 출처", "판정")] + rows,
          [3.0, 1.0, 1.6, 3.2, 3.3], size=12, row_h=0.5, accent_col=4)
    text(s, 0.6, 6.0, 12.1, 1.0, [
        "임계값 미만이면 Tool 결과를 '문서 근거 없음'으로 반환하고, Agent는 추측 답변을 하지 않음",
        ("보정 근거: 샘플 매뉴얼 7청크 기준 관련 질의 최소와 무관 질의 최대의 중간값. 실문서 색인 후 재보정 필요",
         {"color": MUTED, "size": 12}),
    ], size=13, space_after=4)


def s_tools(prs):
    s = new_slide(prs, "Tool 명세", "모든 Tool은 (LLM용 텍스트, UI용 artifact)를 함께 반환 — content_and_artifact")
    rows = [
        ("Tool", "입력", "호출 조건", "LLM에 주는 결과", "UI 표시(artifact)"),
        ("rag_search", "query", "재배 방법·환경·방제·문서 요약", "구절 + [번호]·출처·유사도", "검색 근거 표"),
        ("analyze_mushroom_image", "image_path,\nquestion", "[첨부 이미지 경로]가 있는 질문", "생육 단계·병해충·신뢰도·근거", "구조화 JSON"),
        ("write_farm_diary", "date, notes", "\"일지 써줘\" 등", "저장 경로 + 일지 본문", "저장 경로"),
        ("get_mushroom_price", "period\n(1w~all)", "가격·시세 질문", "도매·소매 기간 요약 + 출처", "요약 표"),
    ]
    table(s, 0.6, 1.75, 12.1, rows, [2.6, 1.6, 2.8, 2.9, 2.2], size=12, row_h=0.72)
    rows2 = [
        ("안전·제한", "내용"),
        ("이미지 경로", "outputs/uploads, data/raw/images, data/sample/images만 허용"),
        ("반복 제한", "최대 6회, 마지막 회차는 Tool 없이 최종 답변"),
        ("오류 처리", "429·연결 오류 지수 백오프 2회 → 'API 한도 초과 또는 모델 서버 미연결'"),
    ]
    table(s, 0.6, 5.55, 12.1, rows2, [2.6, 9.5], size=12, row_h=0.36)


def s_chat_detail(prs):
    s = new_slide(prs, "AI 상담 ① 문서 질의 → rag_search", "Agent 응답과 Tool 호출 이력(입력·검색 근거 표·결과 요약)")
    fit_shot(s, "31_chat_rag", 0.6, 1.65, 6.0, 5.1)
    fit_shot(s, "35_chat_trace_rag", 6.85, 1.65, 5.85, 4.4)
    text(s, 6.85, 6.35, 5.85, 0.6, "mock 응답은 Tool 결과를 그대로 묶은 고정 형식이며 LLM 요약이 아님",
         size=12, color=ACCENT)


def s_chat_tools(prs):
    s = new_slide(prs, "AI 상담 ② 질의 유형별 Tool 선택", "같은 대화창에서 가격·이미지·일지 요청을 각각 다른 Tool로 처리")
    items = [("32_chat_price", "\"최근 가격 알려줘\" → get_mushroom_price(1m)"),
             ("33_chat_image", "이미지 첨부 + \"상태 어때?\" → analyze_mushroom_image"),
             ("34_chat_diary", "\"오늘 일지 써줘. 온도 16도…\" → write_farm_diary")]
    for i, (stem, cap) in enumerate(items):
        x = 0.6 + i * 4.1
        text(s, x, 1.62, 3.9, 0.55, cap, size=12, bold=True, color=FOREST)
        fit_shot(s, stem, x, 2.25, 3.9, 4.5)


def s_diary2(prs):
    s = new_slide(prs, "영농일지 (탭 4 · write_farm_diary)", "src/diary/writer.py · outputs/diary/YYYY-MM-DD.md 저장 + 다운로드")
    fit_shot(s, "04_farm_diary", 0.6, 1.65, 4.6, 5.1)
    fit_shot(s, "41_diary_markdown", 5.4, 1.65, 3.6, 5.1)
    rows = [("양식 항목", "채우는 값"), ("1 날짜", "입력 날짜"), ("2 재배사 환경", "온도 ℃ · 습도 % · CO2 ppm"),
            ("3 생육 상태", "선택한 분석 결과"), ("4 병해충", "선택한 분석 결과 + 근거"),
            ("5 작업 내용", "작업 메모"), ("6 익일 계획", "LLM 생성(입력 외 수치 금지)")]
    table(s, 9.2, 1.65, 3.5, rows, [1.35, 2.15], size=11, row_h=0.5)
    text(s, 9.2, 5.35, 3.5, 1.3, ["미입력 환경값은 '미입력'으로 표기", "모델 오류 시 익일 계획만 실패 문구로 대체"],
         size=12, bullet=True, color=MUTED)


def s_verify2(prs, summary, test_rows):
    s = new_slide(prs, "검증 결과", "LLM_PROVIDER=mock (키 없이) · 빌드 시 pytest 실제 실행")
    rect(s, 0.6, 1.7, 3.9, 1.9, fill=FOREST)
    text(s, 0.8, 1.82, 3.6, 0.4, "pytest 전체", size=14, color=TINT)
    text(s, 0.8, 2.3, 3.6, 1.2, summary, size=22, bold=True, color=WHITE)
    rows = [("테스트 파일", "검증 내용", "통과", "실패", "skip")] + [
        (m, d, str(p), str(f), str(k)) for m, d, p, f, k in test_rows]
    if len(rows) == 1:
        rows.append(("—", "테스트 미실행", "-", "-", "-"))
    table(s, 4.8, 1.7, 7.9, rows, [1.5, 4.3, 0.7, 0.7, 0.7], size=11, row_h=0.42)
    rows2 = [
        ("항목", "상태", "비고"),
        ("실 Embedding (e5-small)", "검증 완료", "RAG 실측 슬라이드 참고"),
        ("Gemini·Ollama 실제 호출", "미검증", "API 키·모델 서버 없음 → mock으로 대체 검증"),
        ("VLM 판독 정확도", "미측정", "라벨링 이미지셋 필요"),
        ("KAMIS Open API", "미검증", "파라미터 [미확인], 접속 불가"),
    ]
    table(s, 0.6, 4.75, 12.1, rows2, [3.4, 1.8, 6.9], size=12, row_h=0.42, accent_col=1)


def s_limits(prs):
    s = new_slide(prs, "한계와 개선점")
    rows = [
        ("데이터", "가격·생산 통계와 이미지가 샘플/더미", "KAMIS·KOSIS 파일과 실제 재배 사진 투입"),
        ("Vision", "Zero-shot VLM, 정량 정확도 미측정", "라벨링 이미지셋 구축 후 정확도 평가, 필요 시 ViT/YOLO 미세조정"),
        ("RAG", "임계값 0.855는 샘플 문서 7청크로만 보정", "공식 매뉴얼 색인 후 질의셋으로 임계값 재보정"),
        ("Agent", "mock은 키워드 규칙이라 복합 질의 미처리", "실모델로 다중 Tool 연쇄 호출 평가"),
        ("가격 API", "KAMIS 파라미터·품목코드 [미확인]", "공식 문서 확인 후 품목코드 .env 설정"),
    ]
    for i, (k, a, b) in enumerate(rows):
        y = 1.5 + i * 1.08
        rect(s, 0.6, y, 2.0, 0.88, fill=FOREST)
        text(s, 0.6, y, 2.0, 0.88, k, size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        rect(s, 2.75, y, 4.85, 0.88, fill=TINT)
        text(s, 2.9, y, 4.6, 0.88, a, size=14, anchor=MSO_ANCHOR.MIDDLE)
        rect(s, 7.75, y, 4.95, 0.88, fill=WHITE, line=MOSS)
        text(s, 7.9, y, 4.7, 0.88, "→ " + b, size=14, color=FOREST, anchor=MSO_ANCHOR.MIDDLE)


def build(skip_tests: bool = False) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    summary, test_rows = run_tests(skip_tests)
    rag = measure_rag()
    s_title(prs)
    s_overview(prs)
    s_requirements(prs, test_rows)
    s_stack(prs)
    s_data(prs)
    s_arch(prs)
    s_dataflow(prs)
    s_dashboard(prs)
    s_dashboard_detail1(prs)
    s_dashboard_detail2(prs)
    s_vision(prs)
    s_vision_flow(prs)
    s_rag(prs)
    s_rag_eval(prs, rag)
    s_agent(prs)
    s_tools(prs)
    s_chat_detail(prs)
    s_chat_tools(prs)
    s_diary2(prs)
    s_demo(prs)
    s_verify2(prs, summary, test_rows)
    s_limits(prs)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"saved {OUT.relative_to(ROOT)} ({len(prs.slides)} slides) · pytest: {summary} · RAG: {rag[0]}")
    return OUT


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--skip-tests", action="store_true")
    build(ap.parse_args().skip_tests)
